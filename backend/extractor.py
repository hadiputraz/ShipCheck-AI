from pathlib import Path

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

import pytesseract
from pdf2image import convert_from_path


# ============================================================
# CONFIGURATION
# ============================================================

TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

POPPLER_PATH = (
    r"C:\Users\HADIPUTRA\Downloads"
    r"\Release-26.09.0-0\poppler-26.09.0\Library\bin"
)

pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


# ============================================================
# TXT
# ============================================================

def extract_txt(path: Path) -> str:
    return path.read_text(
        encoding="utf-8",
        errors="ignore"
    )


# ============================================================
# PDF - NORMAL TEXT EXTRACTION
# ============================================================

def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages)


# ============================================================
# PDF - OCR FALLBACK
# ============================================================

def extract_pdf_ocr(path: Path) -> str:
    print(f"[OCR] Scanned PDF detected: {path.name}")

    try:
        images = convert_from_path(
            str(path),
            dpi=300,
            poppler_path=POPPLER_PATH
        )
    except Exception as e:
        raise RuntimeError(
            f"PDF to image conversion failed: {e}"
        )

    pages = []

    for page_number, image in enumerate(images, start=1):

        print(
            f"[OCR] Processing page "
            f"{page_number}/{len(images)}..."
        )

        text = pytesseract.image_to_string(
            image,
            lang="eng"
        )

        if text:
            pages.append(text)

    return "\n".join(pages)


# ============================================================
# PDF
# ============================================================

def extract_pdf(path: Path) -> str:

    # First attempt normal PDF text extraction
    text = extract_pdf_text(path)

    # If text exists, use it
    if text and text.strip():
        return text

    # Otherwise use OCR
    print(
        f"[PDF] No selectable text found in "
        f"{path.name}. Falling back to OCR."
    )

    ocr_text = extract_pdf_ocr(path)

    return ocr_text


# ============================================================
# DOCX
# ============================================================

def extract_docx(path: Path) -> str:

    document = Document(str(path))

    paragraphs = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    # Extract tables
    for table in document.tables:

        for row in table.rows:

            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(values):
                paragraphs.append(
                    " | ".join(values)
                )

    return "\n".join(paragraphs)


# ============================================================
# XLSX
# ============================================================

def extract_xlsx(path: Path) -> str:

    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True
    )

    lines = []

    for sheet in workbook.worksheets:

        lines.append(
            f"[SHEET: {sheet.title}]"
        )

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                if value is not None:

                    values.append(
                        str(value).strip()
                    )

            if values:

                lines.append(
                    " | ".join(values)
                )

    return "\n".join(lines)


# ============================================================
# MAIN DOCUMENT EXTRACTOR
# ============================================================

def extract_document(path: str) -> str:

    file_path = Path(path)

    if not file_path.exists():

        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = file_path.suffix.lower()

    if extension == ".txt":

        return extract_txt(file_path)

    if extension == ".pdf":

        return extract_pdf(file_path)

    if extension == ".docx":

        return extract_docx(file_path)

    if extension == ".xlsx":

        return extract_xlsx(file_path)

    raise ValueError(
        f"Unsupported document format: {extension}"
    )


# ============================================================
# COMMAND LINE TEST
# ============================================================

if __name__ == "__main__":

    import sys

    if len(sys.argv) != 2:

        print("Usage:")
        print(
            "python backend/extractor.py "
            "<document_path>"
        )

        sys.exit(1)

    document_path = sys.argv[1]

    try:

        text = extract_document(
            document_path
        )

        print("=" * 60)
        print("EXTRACTED DOCUMENT TEXT")
        print("=" * 60)

        print(text)

    except Exception as e:

        print("=" * 60)
        print("EXTRACTION ERROR")
        print("=" * 60)

        print(str(e))

        sys.exit(1)