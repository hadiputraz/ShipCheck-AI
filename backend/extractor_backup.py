from pathlib import Path

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


def extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages)


def extract_docx(path: Path) -> str:
    document = Document(str(path))

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    # Also extract table contents
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]

            if any(values):
                paragraphs.append(" | ".join(values))

    return "\n".join(paragraphs)


def extract_xlsx(path: Path) -> str:
    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True
    )

    lines = []

    for sheet in workbook.worksheets:

        lines.append(f"[SHEET: {sheet.title}]")

        for row in sheet.iter_rows(values_only=True):

            values = []

            for value in row:
                if value is not None:
                    values.append(str(value).strip())

            if values:
                lines.append(" | ".join(values))

    return "\n".join(lines)


def extract_document(path: str) -> str:
    """
    Extract readable text from a supported document.
    """

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    extension = file_path.suffix.lower()

    if extension == ".txt":
        return extract_txt(file_path)

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".xlsx":
        return extract_xlsx(file_path)

    raise ValueError(
        f"Unsupported document format: {extension}"
    )


if __name__ == "__main__":

    import sys

    if len(sys.argv) != 2:
        print("Usage:")
        print("python backend/extractor.py <document_path>")
        sys.exit(1)

    document_path = sys.argv[1]

    text = extract_document(document_path)

    print("=" * 60)
    print("EXTRACTED DOCUMENT TEXT")
    print("=" * 60)
    print(text)